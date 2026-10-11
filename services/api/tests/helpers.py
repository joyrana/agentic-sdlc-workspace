"""Test helpers: in-memory DOCX builders and sample paths."""

from __future__ import annotations

import io
import zipfile
from pathlib import Path

import docx

REPO = Path(__file__).resolve().parents[3]
SAMPLE_BRD = REPO / "samples" / "documents" / "task-title-validation-brd.docx"
SAMPLE_TRANSCRIPT = REPO / "samples" / "documents" / "kickoff-meeting-transcript.txt"


def make_docx(*paragraphs: str | tuple[str, str]) -> bytes:
    """Each item is text (Normal) or (style, text), e.g. ("Heading 1", "Scope")."""
    d = docx.Document()
    for p in paragraphs:
        if isinstance(p, tuple):
            style, text = p
            if style.startswith("Heading"):
                d.add_heading(text, level=int(style.split()[-1]))
            else:
                d.add_paragraph(text, style=style)
        else:
            d.add_paragraph(p)
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def make_zip(entries: dict[str, bytes], compression: int = zipfile.ZIP_DEFLATED) -> bytes:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", compression=compression) as zf:
        for name, data in entries.items():
            zf.writestr(name, data)
    return buf.getvalue()
