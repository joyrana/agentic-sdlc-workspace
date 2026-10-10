"""Generate the sample BRD (.docx) used by the demo and tests.

SAMPLE DATA. The document describes a small change to the sample repository
`samples/repositories/sample-task-service`. Run: python samples/generate_samples.py
"""

from __future__ import annotations

from pathlib import Path

import docx
from docx.shared import Pt

OUT = Path(__file__).resolve().parent / "documents"


def build_brd(path: Path) -> None:
    d = docx.Document()
    d.core_properties.title = "Task Service: Title Validation BRD (sample)"
    d.core_properties.author = "Sample Product Team"
    style = d.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    d.add_heading("Task Service: Title Validation", level=0)
    d.add_paragraph(
        "Sample business requirements document shipped with the Agentic SDLC Workspace. "
        "It describes a small, safe change to the sample task service."
    )

    d.add_heading("1 Purpose", level=1)
    d.add_paragraph(
        "Support agents report tasks with blank or extremely long titles. These break the task list and "
        "exports. This change makes task titles predictable."
    )

    d.add_heading("2 Scope", level=1)
    d.add_paragraph("In scope: validation of task titles when a task is created.", style="List Bullet")
    d.add_paragraph("Out of scope: editing existing tasks and descriptions.", style="List Bullet")

    d.add_heading("3 Functional requirements", level=1)
    table = d.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    for i, h in enumerate(["ID", "Requirement", "Priority", "Acceptance criteria"]):
        hdr[i].text = h
    rows = [
        (
            "FR-001",
            "The system shall reject task titles longer than 120 characters with a validation error.",
            "Must",
            "Given a title of 121 characters, when the task is created, then a validation error is raised; "
            "Given a title of exactly 120 characters, when the task is created, then the task is created",
        ),
        (
            "FR-002",
            "The system shall trim leading and trailing whitespace from task titles before saving.",
            "Must",
            "Given the title '  Buy milk  ', when the task is created, then the stored title is 'Buy milk'",
        ),
        (
            "FR-003",
            "The system shall reject titles that are empty after trimming (see FR-002).",
            "Must",
            "Given a title containing only spaces, when the task is created, then a validation error is raised",
        ),
    ]
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = value

    d.add_heading("4 Non-functional requirements", level=1)
    d.add_paragraph("NFR-001: Task creation should be fast.")
    d.add_paragraph(
        "NFR-002: Validation errors must include the field name and a human-readable message."
    )
    d.add_heading("Acceptance criteria", level=2)
    d.add_paragraph(
        "Given an invalid title, when the task is created, then the error names the field 'title'.",
        style="List Bullet",
    )

    d.add_heading("5 Business rules", level=1)
    d.add_paragraph("BR-001: If a title is rejected, then no task is stored.")

    d.add_heading("6 Open items", level=1)
    d.add_paragraph("The maximum description length is TBD.")
    d.save(str(path))


if __name__ == "__main__":
    OUT.mkdir(parents=True, exist_ok=True)
    build_brd(OUT / "task-title-validation-brd.docx")
    print("wrote", OUT / "task-title-validation-brd.docx")
